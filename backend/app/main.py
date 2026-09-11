from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from typing import Optional
import os
import uuid
import logging
import uvicorn
import io
import pypdf
import docx
from docx.oxml.ns import qn

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request, Header, Depends
from fastapi.responses import JSONResponse
from slowapi import Limiter
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
from app.database import init_db, SessionLocal, AnalysisResult
from app.tasks import process_resume_analysis

load_dotenv()

# Logger configuration
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("API")

app = FastAPI(title="AgenticHire API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",  # For local development
        "https://agentic-hire-lf9py5dca-elite-juniors-projects.vercel.app"
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Rate limiting ---
# Backed by the same Redis instance already used as the Celery broker (see
# app/tasks.py) so the limit is shared correctly across every API process --
# an in-memory limiter would reset per-process and not actually coordinate
# across multiple Uvicorn/Railway instances.
# Caveat: get_remote_address reads request.client.host, which is only the
# *real* client IP if the ASGI server is told to trust the platform's
# reverse-proxy forwarded headers (e.g. `uvicorn --proxy-headers`). Without
# that, every request behind Railway's proxy can appear to share one IP,
# making the limit effectively global rather than per-visitor -- still a
# real cap on total abuse, just not perfectly per-user in that setup.
REDIS_URL = os.getenv("REDIS_URL", "redis://localhost:6379/0")
limiter = Limiter(key_func=get_remote_address, storage_uri=REDIS_URL)
app.state.limiter = limiter


@app.exception_handler(RateLimitExceeded)
async def rate_limit_handler(request: Request, exc: RateLimitExceeded):
    # Matches the {"detail": "..."} shape FastAPI's own HTTPException uses,
    # so the frontend's existing error-popup handling (see
    # frontend/app/api/analyze/route.ts) shows this cleanly too.
    return JSONResponse(
        status_code=429,
        content={"detail": "⏳ Too many requests. Please wait a minute and try again."},
    )


# --- Internal API secret ---
# Only the Next.js server (frontend/app/api/*/route.ts -- never the browser
# bundle) knows this value, sent as a header on every proxied request. This
# doesn't gate real end users at all (the site stays login-free) -- it only
# stops someone calling this API directly (curl/Postman/a script) bypassing
# the frontend entirely.
#
# Fails OPEN (unprotected, with a loud warning) if the env var isn't set at
# all, rather than rejecting every request outright -- so local dev/testing
# without it configured keeps working exactly as it did before this change.
# In production this variable MUST be set (Railway env vars) for this to
# actually protect anything; an unset var there is a silent no-op, not an
# error, so it's easy to forget -- hence the startup warning below.
INTERNAL_API_SECRET = os.getenv("INTERNAL_API_SECRET")

if not INTERNAL_API_SECRET:
    logger.warning(
        "⚠️ INTERNAL_API_SECRET is not set — /analyze and /status are "
        "UNPROTECTED right now. Fine for local dev, but this MUST be set "
        "in production for this check to do anything."
    )


async def verify_internal_secret(x_internal_api_key: Optional[str] = Header(default=None)):
    if INTERNAL_API_SECRET and x_internal_api_key != INTERNAL_API_SECRET:
        raise HTTPException(status_code=401, detail="Unauthorized")


@app.on_event("startup")
def startup():
    init_db()


# --- One-page-resume enforcement ---
MAX_RESUME_PAGES = 1


def count_pages(filename: str, content: bytes) -> Optional[int]:
    """
    Best-effort page count, used to reject resumes over MAX_RESUME_PAGES.

    - PDF: exact and reliable -- pypdf reports the real number of page
      objects in the file, regardless of whether those pages contain a text
      layer or are scanned images (this runs before the OCR fallback, so a
      multi-page scanned resume gets rejected here instead of wasting a
      Gemini OCR call on it).
    - DOCX: NOT reliable. Word documents are reflowable -- there is no
      stored "page count" until something actually renders/paginates the
      file (fonts, margins, page size all affect it), which python-docx
      does not do. This only counts explicit manual page breaks
      (`<w:br w:type="page"/>`), so it can under-count (a doc that
      naturally overflows to page 2 without one won't be caught) but never
      over-counts -- if it reports >1, that's real evidence, not a guess.
    - TXT: no notion of a "page" at all (no layout/formatting). Not
      enforced.
    """
    try:
        if filename.lower().endswith(".pdf"):
            return len(pypdf.PdfReader(io.BytesIO(content)).pages)

        if filename.lower().endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            page_breaks = sum(
                1
                for br in doc.element.body.iter(qn("w:br"))
                if br.get(qn("w:type")) == "page"
            )
            return page_breaks + 1

    except Exception as e:
        logger.error(f"Could not determine page count: {e}")
        return None

    return None


# --- Extraction functions inside the API ---
async def extract_text_from_upload(file: UploadFile) -> tuple[str, bytes]:
    # Read the file into memory (Bytes)
    content = await file.read()
    file_obj = io.BytesIO(content)
    text = ""

    try:
        filename = file.filename.lower()
        if filename.endswith('.pdf'):
            reader = pypdf.PdfReader(file_obj)
            for page in reader.pages:
                text += page.extract_text() + "\n"

        elif filename.endswith('.docx'):
            doc = docx.Document(file_obj)
            for para in doc.paragraphs:
                text += para.text + "\n"

        elif filename.endswith('.txt'):
            text = content.decode('utf-8')

    except Exception as e:
        logger.error(f"Error extracting text: {e}")
        return "", content

    return text, content


# ---------------------------------

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB -- generous for a 1-page resume


@app.post("/analyze", dependencies=[Depends(verify_internal_secret)])
@limiter.limit("10/minute")
async def start_analysis(request: Request, job_description: str = Form(...), file: UploadFile = File(...)):
    session_id = str(uuid.uuid4())
    logger.info(f"🔵 NEW REQUEST: {session_id}")

    # 0a. Reject oversized uploads outright -- before anything else touches
    # the file (page-count parsing, text extraction, and especially the OCR
    # fallback, which would otherwise ship a huge payload to Gemini).
    raw_bytes_preview = await file.read()
    await file.seek(0)
    if len(raw_bytes_preview) > MAX_FILE_SIZE_BYTES:
        logger.info(f"🚫 REJECTED {session_id}: file is {len(raw_bytes_preview)} bytes (limit: {MAX_FILE_SIZE_BYTES})")
        raise HTTPException(
            status_code=400,
            detail=f"⛔ REJECTED: File too large ({len(raw_bytes_preview) / 1_048_576:.1f} MB). Max size is {MAX_FILE_SIZE_BYTES // 1_048_576} MB.",
        )

    # 0b. Reject multi-page resumes outright, before spending any effort on
    # extraction/OCR. A resume should fit on one page -- see count_pages()
    # for what's actually reliable per file type (PDF: exact; DOCX:
    # best-effort; TXT: unenforced).
    page_count = count_pages(file.filename, raw_bytes_preview)
    if page_count is not None and page_count > MAX_RESUME_PAGES:
        logger.info(f"🚫 REJECTED {session_id}: resume has {page_count} pages (limit: {MAX_RESUME_PAGES})")
        raise HTTPException(
            status_code=400,
            detail=(
                f"⛔ REJECTED: Your resume has {page_count} pages. "
                "A professional resume MUST fit on a SINGLE page. "
                "Please condense it to one page and try again."
            ),
        )

    # 1. Extract the text in memory (without saving a file)
    resume_text, raw_bytes = await extract_text_from_upload(file)

    # 1b. OCR fallback: the PDF may be a scanned/photographed resume with no
    # extractable text layer -- hand the raw bytes to Gemini directly instead
    # of failing outright. Only worth trying for PDFs (scanned .docx/.txt
    # don't really occur in practice).
    if not resume_text.strip() and file.filename.lower().endswith(".pdf"):
        logger.info(f"🔎 No embedded text found for {session_id} — falling back to OCR via Gemini")
        try:
            from agno.media import File as AgnoFile
            from app.agents import resume_ocr_agent

            ocr_file = AgnoFile(content=raw_bytes, mime_type="application/pdf")
            ocr_result = resume_ocr_agent.run("Transcribe this resume.", files=[ocr_file])
            resume_text = str(ocr_result.content or "").strip()
        except Exception as e:
            logger.error(f"OCR fallback failed: {e}")

    if not resume_text.strip():
        raise HTTPException(status_code=400, detail="Could not extract text from file")

    # 2. Save to DB
    db = SessionLocal()
    try:
        new_analysis = AnalysisResult(
            session_id=session_id,
            status="pending",
            job_description=job_description
        )
        db.add(new_analysis)
        db.commit()
    except Exception as e:
        logger.error(f"DB Error: {e}")
        return {"error": "Database error"}
    finally:
        db.close()

    # 3. Send the text (not the file) to Worker
    try:
        # Sending the text as an argument
        process_resume_analysis.delay(session_id, resume_text, job_description)
        logger.info("🟢 Task sent to Celery")
    except Exception as e:
        logger.error(f"Celery Error: {e}")
        return {"error": "Failed to queue task"}

    return {"session_id": session_id, "status": "processing"}


@app.get("/status/{session_id}", dependencies=[Depends(verify_internal_secret)])
async def get_status(session_id: str):
    db = SessionLocal()
    try:
        result = db.query(AnalysisResult).filter(AnalysisResult.session_id == session_id).first()
        if not result:
            return {"error": "Not found"}

        return {
            "status": result.status,
            "result": result.result_metadata,
            "formatted_text": result.result_text
        }
    finally:
        db.close()


if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)