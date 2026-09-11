"""
Pure text/metadata extraction for uploaded resume files -- page counting and
plain-text extraction from PDF/DOCX/TXT bytes.

Deliberately has ZERO dependency on app.tasks/app.agents/app.database:
importing those triggers Langfuse's Strict Mode prompt fetch at module load
time (real GOOGLE_API_KEY/LANGFUSE_* credentials + network required -- see
the observability-langfuse skill), which would make this module
unimportable in CI without secrets. Keeping it dependency-free lets its
logic actually be unit-tested (see backend/tests/test_extraction.py) with no
secrets and no external services -- pulled out of main.py for exactly that
reason.
"""
import io
import logging
from typing import Optional

import pypdf
import docx
from docx.oxml.ns import qn

logger = logging.getLogger("API")

MAX_RESUME_PAGES = 1
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB -- generous for a 1-page resume


def count_pages(filename: str, content: bytes) -> Optional[int]:
    """
    Best-effort page count, used to reject resumes over MAX_RESUME_PAGES.

    - PDF: exact and reliable -- pypdf reports the real number of page
      objects in the file, regardless of whether those pages contain a text
      layer or are scanned images (this runs before the OCR fallback, so a
      multi-page scanned resume gets rejected here instead of wasting a
      Gemini OCR call on it).
    - DOCX: NOT reliable. Word documents are reflowable -- there is no
      stored "page count" until something actually renders/paginates the
      file (fonts, margins, page size all affect it), which python-docx
      does not do. This only counts explicit manual page breaks
      (`<w:br w:type="page"/>`), so it can under-count (a doc that
      naturally overflows to page 2 without one won't be caught) but never
      over-counts -- if it reports >1, that's real evidence, not a guess.
    - TXT: no notion of a "page" at all (no layout/formatting). Not
      enforced.
    """
    try:
        if filename.lower().endswith(".pdf"):
            return len(pypdf.PdfReader(io.BytesIO(content)).pages)

        if filename.lower().endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            page_breaks = sum(
                1
                for br in doc.element.body.iter(qn("w:br"))
                if br.get(qn("w:type")) == "page"
            )
            return page_breaks + 1

    except Exception as e:
        logger.error(f"Could not determine page count: {e}")
        return None

    return None


def extract_text_from_content(filename: str, content: bytes) -> str:
    """
    Extract plain text from PDF/DOCX/TXT bytes.

    Returns "" on failure OR for an unsupported extension -- callers (see
    main.py) treat an empty string as "try the OCR fallback, then give up
    with a 400" uniformly, regardless of which of those two cases it was.
    """
    file_obj = io.BytesIO(content)
    text = ""

    try:
        lower = filename.lower()
        if lower.endswith(".pdf"):
            reader = pypdf.PdfReader(file_obj)
            for page in reader.pages:
                text += page.extract_text() + "\n"

        elif lower.endswith(".docx"):
            doc = docx.Document(file_obj)
            for para in doc.paragraphs:
                text += para.text + "\n"

        elif lower.endswith(".txt"):
            text = content.decode("utf-8")

    except Exception as e:
        logger.error(f"Error extracting text: {e}")
        return ""

    return text
